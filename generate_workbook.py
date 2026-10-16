"""
generate_workbook.py
Builds HR_Dashboard.xlsx: an HR operations workbook with data validation,
formula-driven checks and a KPI dashboard.

Run:  python generate_workbook.py
"""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")   # light yellow = type here
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BASE = Font(name=FONT, size=10)
BLUE = Font(name=FONT, size=10, color="0000FF")        # hard-coded inputs
BOLD = Font(name=FONT, size=10, bold=True)
thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LAST = 60  # formulas / validation are pre-filled to this row so new rows work

wb = Workbook()


def make_sheet(title, headers, widths):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    return ws


def style_body(ws, ncols, first=2, last=LAST):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if cell.font == Font():  # untouched default
                cell.font = BASE


def add_list_validation(ws, rng, options, msg):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Invalid entry", error=msg)
    ws.add_data_validation(dv)
    dv.add(rng)


def add_date_validation(ws, rng):
    dv = DataValidation(type="date", operator="greaterThanOrEqual",
                        formula1="36526", allow_blank=True,  # 01-Jan-2000
                        showErrorMessage=True, errorTitle="Invalid date",
                        error="Enter a valid date (on or after 01-Jan-2000).")
    ws.add_data_validation(dv)
    dv.add(rng)


def add_nonneg_validation(ws, rng):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual",
                        formula1="0", allow_blank=True, showErrorMessage=True,
                        errorTitle="Invalid amount", error="Enter a number that is 0 or more.")
    ws.add_data_validation(dv)
    dv.add(rng)


wb.remove(wb.active)

# ---------------------------------------------------------------- README
rd = wb.create_sheet("README")
rd.column_dimensions["A"].width = 110
lines = [
    ("HR Analytics & Employee Records Dashboard", Font(name=FONT, size=14, bold=True, color="1F3A5F")),
    ("", BASE),
    ("HOW TO USE", BOLD),
    ("1. Yellow cells are inputs - type or paste your data there. White cells with formulas calculate automatically.", BASE),
    ("2. Sheets: Employees, Candidates, Onboarding, Attendance, Payroll_Input -> all feed the Dashboard sheet.", BASE),
    ("3. Drop-down lists and number/date rules stop invalid entries (e.g. a wrong Status or a negative salary).", BASE),
    ("4. Payroll_Input column I flags problems: missing data, unknown employee, duplicate, deductions above gross, low basic %.", BASE),
    ("5. Formulas are pre-filled to row 60 on each sheet, so you can add rows without copying formulas.", BASE),
    ("", BASE),
    ("ASSUMPTIONS", BOLD),
    ("- All names, amounts and dates are SAMPLE data created for demonstration. Replace with real data.", BASE),
    ("- Minimum Basic % of gross is a company-policy setting (Settings!B2, default 40%). Change it to match your policy.", BASE),
    ("- Last 4 rows of Payroll_Input are deliberate test cases (see Note column) to show the validation. Delete before real use.", BASE),
    ("- Attrition % = Resigned / (Active + Resigned) in this file's data. It is a simple point-in-time ratio, not an annualised rate.", BASE),
    ("- Keep personal data (PAN, bank account numbers, Aadhaar) OUT of this file; use your HRMS/payroll system for those.", BASE),
]
for i, (t, f) in enumerate(lines, start=1):
    rd.cell(row=i, column=1, value=t).font = f
    rd.cell(row=i, column=1).alignment = Alignment(wrap_text=True)

st = wb.create_sheet("Settings")
st.column_dimensions["A"].width = 40
st.column_dimensions["B"].width = 14
st["A1"], st["B1"] = "Setting", "Value"
for c in ("A1", "B1"):
    st[c].font, st[c].fill = HDR_FONT, HDR_FILL
st["A2"], st["B2"] = "Minimum Basic as % of Gross (policy)", 0.40
st["A2"].font = BASE
st["B2"].font = BLUE
st["B2"].number_format = "0%"
st["B2"].fill = INPUT_FILL
st["A3"] = "Assumption: 40% is a placeholder. Set per your company's salary-structure policy."
st["A3"].font = Font(name=FONT, size=9, italic=True)

# ---------------------------------------------------------------- Employees
emp = make_sheet("Employees",
                 ["Emp ID", "Employee Name", "Department", "Designation", "Join Date",
                  "Status", "Exit Date", "Monthly CTC (INR)"],
                 [11, 22, 16, 24, 13, 12, 13, 18])
employees = [
    ("EMP001", "Aarav Menon", "HR", "HR Executive", date(2024, 6, 3), "Active", None, 32000),
    ("EMP002", "Divya Raman", "Sales", "Sales Associate", date(2024, 8, 12), "Active", None, 28000),
    ("EMP003", "Karthik S", "Engineering", "Software Engineer", date(2024, 9, 2), "Active", None, 55000),
    ("EMP004", "Meera Nair", "Finance", "Accounts Executive", date(2025, 1, 6), "Active", None, 36000),
    ("EMP005", "Rohan Das", "Operations", "Operations Associate", date(2025, 2, 17), "Resigned", date(2026, 5, 31), 30000),
    ("EMP006", "Sneha Iyer", "HR", "Payroll Associate", date(2025, 4, 7), "Active", None, 34000),
    ("EMP007", "Vikram Patel", "Sales", "Sales Executive", date(2025, 7, 1), "Resigned", date(2026, 3, 31), 31000),
    ("EMP008", "Anjali Rao", "Engineering", "QA Engineer", date(2026, 1, 12), "Active", None, 42000),
    ("EMP009", "Harish K", "Operations", "Operations Associate", date(2026, 2, 2), "Active", None, 29000),
    ("EMP010", "Pooja Verma", "Finance", "Finance Analyst", date(2026, 3, 16), "Active", None, 45000),
    ("EMP011", "Naveen T", "Sales", "Sales Associate", date(2026, 6, 1), "Active", None, 27000),
    ("EMP012", "Lakshmi P", "HR", "Talent Acquisition Associate", date(2026, 7, 6), "Active", None, 33000),
]
for r, row in enumerate(employees, start=2):
    for c, v in enumerate(row, start=1):
        cell = emp.cell(row=r, column=c, value=v)
        cell.font = BLUE
for r in range(2, LAST + 1):
    for c in range(1, 9):
        emp.cell(row=r, column=c).fill = INPUT_FILL
    emp.cell(row=r, column=5).number_format = "dd-mmm-yyyy"
    emp.cell(row=r, column=7).number_format = "dd-mmm-yyyy"
    emp.cell(row=r, column=8).number_format = "#,##0"
style_body(emp, 8)
add_list_validation(emp, f"F2:F{LAST}", ["Active", "Resigned"], "Status must be Active or Resigned.")
add_date_validation(emp, f"E2:E{LAST}")
add_date_validation(emp, f"G2:G{LAST}")
add_nonneg_validation(emp, f"H2:H{LAST}")
# Flag duplicate Emp IDs
emp.conditional_formatting.add(
    f"A2:A{LAST}",
    FormulaRule(formula=[f'AND(A2<>"",COUNTIF($A$2:$A${LAST},A2)>1)'],
                fill=PatternFill("solid", bgColor="F8CBAD")))

# ---------------------------------------------------------------- Candidates
STAGES = ["Applied", "Screening", "Interview", "Offer", "Joined", "Rejected"]
can = make_sheet("Candidates",
                 ["Candidate ID", "Candidate Name", "Position Applied", "Source",
                  "Stage", "Applied Date", "Recruiter"],
                 [13, 22, 26, 14, 12, 14, 14])
cands = [
    ("C001", "Ishaan Kumar", "HR Executive", "LinkedIn", "Joined", date(2026, 5, 4), "Lakshmi P"),
    ("C002", "Tanya Singh", "Sales Associate", "Naukri", "Joined", date(2026, 5, 10), "Lakshmi P"),
    ("C003", "Arjun V", "Software Engineer", "Referral", "Offer", date(2026, 8, 2), "Aarav Menon"),
    ("C004", "Nisha B", "Payroll Associate", "LinkedIn", "Interview", date(2026, 8, 14), "Lakshmi P"),
    ("C005", "Farhan Ali", "Sales Executive", "Naukri", "Rejected", date(2026, 8, 18), "Aarav Menon"),
    ("C006", "Gayathri M", "Finance Analyst", "Campus", "Screening", date(2026, 9, 1), "Lakshmi P"),
    ("C007", "Ramesh G", "QA Engineer", "Referral", "Interview", date(2026, 9, 3), "Aarav Menon"),
    ("C008", "Swathi R", "HR Executive", "LinkedIn", "Applied", date(2026, 9, 10), "Lakshmi P"),
    ("C009", "Dev Malhotra", "Operations Associate", "Naukri", "Rejected", date(2026, 9, 12), "Aarav Menon"),
    ("C010", "Preethi S", "Sales Associate", "Campus", "Screening", date(2026, 9, 15), "Lakshmi P"),
    ("C011", "Yusuf Khan", "Software Engineer", "LinkedIn", "Applied", date(2026, 9, 20), "Aarav Menon"),
    ("C012", "Bhavna J", "Talent Acquisition Associate", "Referral", "Offer", date(2026, 9, 22), "Lakshmi P"),
    ("C013", "Siddharth N", "Finance Analyst", "Naukri", "Applied", date(2026, 9, 25), "Lakshmi P"),
    ("C014", "Kavya L", "HR Executive", "Campus", "Interview", date(2026, 9, 27), "Aarav Menon"),
    ("C015", "Mohan Raj", "Operations Associate", "LinkedIn", "Screening", date(2026, 10, 1), "Lakshmi P"),
]
for r, row in enumerate(cands, start=2):
    for c, v in enumerate(row, start=1):
        can.cell(row=r, column=c, value=v).font = BLUE
for r in range(2, LAST + 1):
    for c in range(1, 8):
        can.cell(row=r, column=c).fill = INPUT_FILL
    can.cell(row=r, column=6).number_format = "dd-mmm-yyyy"
style_body(can, 7)
add_list_validation(can, f"E2:E{LAST}", STAGES, "Pick a valid stage from the list.")
add_list_validation(can, f"D2:D{LAST}", ["LinkedIn", "Naukri", "Referral", "Campus", "Other"], "Pick a source from the list.")
add_date_validation(can, f"F2:F{LAST}")

# ---------------------------------------------------------------- Onboarding
TASKS = ["Offer Letter Signed", "Documents Collected", "System/ID Created", "Induction Done", "Policies Acknowledged"]
onb = make_sheet("Onboarding", ["Emp ID", "Employee Name"] + TASKS + ["Completion %"],
                 [11, 22, 14, 16, 16, 14, 16, 14])
onb_rows = [
    ("EMP008", "Anjali Rao", "Yes", "Yes", "Yes", "Yes", "Yes"),
    ("EMP009", "Harish K", "Yes", "Yes", "Yes", "Yes", "Yes"),
    ("EMP010", "Pooja Verma", "Yes", "Yes", "Yes", "Yes", "No"),
    ("EMP011", "Naveen T", "Yes", "Yes", "Yes", "No", "No"),
    ("EMP012", "Lakshmi P", "Yes", "Yes", "No", "No", "No"),
]
for r, row in enumerate(onb_rows, start=2):
    for c, v in enumerate(row, start=1):
        onb.cell(row=r, column=c, value=v).font = BLUE
for r in range(2, LAST + 1):
    for c in range(1, 8):
        onb.cell(row=r, column=c).fill = INPUT_FILL
    onb.cell(row=r, column=8, value=f'=IF(A{r}="","",COUNTIF(C{r}:G{r},"Yes")/COLUMNS(C{r}:G{r}))')
    onb.cell(row=r, column=8).number_format = "0%"
style_body(onb, 8)
add_list_validation(onb, f"C2:G{LAST}", ["Yes", "No"], "Enter Yes or No.")

# ---------------------------------------------------------------- Attendance
att = make_sheet("Attendance",
                 ["Emp ID", "Month", "Working Days", "Days Present", "Days Absent / On Leave", "Attendance %"],
                 [11, 12, 14, 14, 18, 14])
months = [("2026-07", 23), ("2026-08", 21), ("2026-09", 22)]
present_map = {
    "EMP001": (22, 20, 21), "EMP002": (21, 19, 22), "EMP003": (23, 21, 22),
    "EMP004": (20, 20, 21), "EMP006": (22, 21, 22), "EMP008": (23, 18, 20),
    "EMP009": (21, 21, 22), "EMP010": (22, 20, 19), "EMP011": (23, 21, 22),
    "EMP012": (20, 20, 21),
}
r = 2
for eid, pres in present_map.items():
    for (m, wd), p in zip(months, pres):
        att.cell(row=r, column=1, value=eid).font = BLUE
        att.cell(row=r, column=2, value=m).font = BLUE
        att.cell(row=r, column=3, value=wd).font = BLUE
        att.cell(row=r, column=4, value=p).font = BLUE
        r += 1
for r in range(2, LAST + 31):  # leave room for more rows
    for c in range(1, 5):
        att.cell(row=r, column=c).fill = INPUT_FILL
    att.cell(row=r, column=5, value=f'=IF(A{r}="","",C{r}-D{r})')
    att.cell(row=r, column=6, value=f'=IF(A{r}="","",IF(C{r}>0,D{r}/C{r},0))')
    att.cell(row=r, column=6).number_format = "0.0%"
style_body(att, 6, 2, LAST + 30)
dv = DataValidation(type="whole", operator="between", formula1="0", formula2="31",
                    allow_blank=True, showErrorMessage=True, errorTitle="Invalid days",
                    error="Enter a whole number between 0 and 31.")
att.add_data_validation(dv)
dv.add(f"C2:D{LAST + 30}")
# Highlight days present > working days
att.conditional_formatting.add(
    f"D2:D{LAST + 30}",
    FormulaRule(formula=[f'AND(D2<>"",C2<>"",D2>C2)'], fill=PatternFill("solid", bgColor="F8CBAD")))

# ---------------------------------------------------------------- Payroll_Input
pay = make_sheet("Payroll_Input",
                 ["Emp ID", "Month", "Basic (INR)", "HRA (INR)", "Allowances (INR)",
                  "Deductions (INR)", "Gross (INR)", "Net Pay (INR)", "Validation Check", "Note"],
                 [11, 11, 13, 12, 15, 15, 13, 14, 30, 30])
base = {"EMP001": 16000, "EMP002": 14000, "EMP003": 27500, "EMP004": 18000, "EMP006": 17000,
        "EMP008": 21000, "EMP009": 14500, "EMP010": 22500, "EMP011": 13500, "EMP012": 16500}
r = 2
for eid, b in base.items():
    pay.cell(row=r, column=1, value=eid)
    pay.cell(row=r, column=2, value="2026-09")
    pay.cell(row=r, column=3, value=b)
    pay.cell(row=r, column=4, value=round(b * 0.4))
    pay.cell(row=r, column=5, value=round(b * 0.1))
    pay.cell(row=r, column=6, value=round(b * 0.12))
    r += 1
# deliberate test cases
tests = [
    ("EMP003", "2026-09", 27500, 11000, 2750, 12000, "Test: duplicate of row above? (see Check)"),
    ("EMP099", "2026-09", 15000, 6000, 1500, 1800, "Test: unknown employee"),
    ("EMP004", "2026-08", 18000, 7200, 1800, 40000, "Test: deductions exceed gross"),
    ("EMP006", "2026-08", None, 6800, 1700, 2000, "Test: basic missing"),
]
# make first test a true duplicate of EMP003 / 2026-09
for row in tests:
    for c, v in enumerate(row[:6], start=1):
        pay.cell(row=r, column=c, value=v)
    pay.cell(row=r, column=10, value=row[6])
    r += 1
pay["J12"] = "Test: duplicate Emp ID + Month"
for rr in range(2, r):
    for c in range(1, 7):
        pay.cell(row=rr, column=c).font = BLUE
for rr in range(2, LAST + 1):
    for c in range(1, 7):
        pay.cell(row=rr, column=c).fill = INPUT_FILL
    pay.cell(row=rr, column=7, value=f'=IF(A{rr}="","",SUM(C{rr}:E{rr}))')
    pay.cell(row=rr, column=8, value=f'=IF(A{rr}="","",G{rr}-F{rr})')
    pay.cell(row=rr, column=9, value=(
        f'=IF(A{rr}="","",'
        f'IF(OR(B{rr}="",C{rr}="",F{rr}=""),"Missing data",'
        f'IF(COUNTIF(Employees!$A$2:$A${LAST},A{rr})=0,"Unknown employee",'
        f'IF(COUNTIFS($A$2:$A${LAST},A{rr},$B$2:$B${LAST},B{rr})>1,"Duplicate entry",'
        f'IF(F{rr}>G{rr},"Deductions exceed gross",'
        f'IF(C{rr}<Settings!$B$2*G{rr},"Basic below policy %","OK"))))))'))
    for c in (3, 4, 5, 6, 7, 8):
        pay.cell(row=rr, column=c).number_format = "#,##0"
    pay.cell(row=rr, column=10).font = Font(name=FONT, size=9, italic=True)
style_body(pay, 10)
add_nonneg_validation(pay, f"C2:F{LAST}")
pay.conditional_formatting.add(
    f"I2:I{LAST}",
    FormulaRule(formula=['AND(I2<>"",I2<>"OK")'], fill=PatternFill("solid", bgColor="F8CBAD"),
                font=Font(name=FONT, bold=True, color="9C0006")))
pay.conditional_formatting.add(
    f"I2:I{LAST}",
    CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", bgColor="C6EFCE")))
pay["I1"].comment = Comment("Rules: missing data, unknown employee, duplicate (Emp ID + Month), "
                            "deductions > gross, basic < Settings!B2 x gross.", "HR Dashboard")

# ---------------------------------------------------------------- Dashboard
dash = wb.create_sheet("Dashboard", 0)
dash.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [30, 14, 4, 26, 14, 4, 22, 14]):
    dash.column_dimensions[col].width = w
dash["A1"] = "HR Analytics Dashboard"
dash["A1"].font = Font(name=FONT, size=16, bold=True, color="1F3A5F")
dash["A2"] = "All figures are formulas linked to the data sheets. Sample data shown."
dash["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")


def kpi_header(cell, text):
    dash[cell] = text
    dash[cell].font, dash[cell].fill = HDR_FONT, HDR_FILL


def kpi(row, label, formula, fmt="0"):
    dash.cell(row=row, column=1, value=label).font = BASE
    c = dash.cell(row=row, column=2, value=formula)
    c.font, c.number_format = BOLD, fmt
    c.alignment = Alignment(horizontal="right")
    dash.cell(row=row, column=1).border = BORDER
    c.border = BORDER


kpi_header("A4", "Workforce")
kpi_header("B4", "Value")
kpi(5, "Active headcount", f'=COUNTIF(Employees!$F$2:$F${LAST},"Active")')
kpi(6, "Resigned", f'=COUNTIF(Employees!$F$2:$F${LAST},"Resigned")')
kpi(7, "Attrition % (Resigned / total)", '=IFERROR(B6/(B5+B6),0)', "0.0%")
kpi(8, "Avg monthly CTC - active (INR)",
    f'=IFERROR(AVERAGEIFS(Employees!$H$2:$H${LAST},Employees!$F$2:$F${LAST},"Active"),0)', "#,##0")

kpi_header("A10", "Headcount by Department (Active)")
kpi_header("B10", "Count")
depts = ["HR", "Sales", "Engineering", "Finance", "Operations"]
for i, d in enumerate(depts):
    rr = 11 + i
    dash.cell(row=rr, column=1, value=d).font = BLUE
    dash.cell(row=rr, column=1).border = BORDER
    c = dash.cell(row=rr, column=2,
                  value=f'=COUNTIFS(Employees!$C$2:$C${LAST},A{rr},Employees!$F$2:$F${LAST},"Active")')
    c.font, c.border = BOLD, BORDER

kpi_header("D4", "Hiring Pipeline (Stage)")
kpi_header("E4", "Candidates")
for i, s in enumerate(STAGES):
    rr = 5 + i
    dash.cell(row=rr, column=4, value=s).font = BLUE
    dash.cell(row=rr, column=4).border = BORDER
    c = dash.cell(row=rr, column=5, value=f'=COUNTIF(Candidates!$E$2:$E${LAST},D{rr})')
    c.font, c.border = BOLD, BORDER
dash.cell(row=11, column=4, value="Total candidates").font = BASE
dash["E11"] = "=SUM(E5:E10)"
dash["E11"].font = BOLD
dash.cell(row=12, column=4, value="Hire conversion (Joined / Total)").font = BASE
dash["E12"] = "=IFERROR(E9/E11,0)"
dash["E12"].font, dash["E12"].number_format = BOLD, "0.0%"

kpi_header("G4", "Attendance Trend")
kpi_header("H4", "Avg Attendance %")
for i, (m, _) in enumerate(months):
    rr = 5 + i
    dash.cell(row=rr, column=7, value=m).font = BLUE
    dash.cell(row=rr, column=7).border = BORDER
    c = dash.cell(row=rr, column=8,
                  value=f'=IFERROR(AVERAGEIFS(Attendance!$F$2:$F${LAST + 30},Attendance!$B$2:$B${LAST + 30},G{rr}),0)')
    c.font, c.number_format, c.border = BOLD, "0.0%", BORDER
dash.cell(row=9, column=7, value="Overall").font = BASE
dash["H9"] = f"=IFERROR(AVERAGE(Attendance!$F$2:$F${LAST + 30}),0)"
dash["H9"].font, dash["H9"].number_format = BOLD, "0.0%"

kpi_header("G11", "Onboarding & Payroll Quality")
kpi_header("H11", "Value")
dash["G12"] = "Avg onboarding completion"
dash["H12"] = f"=IFERROR(AVERAGE(Onboarding!$H$2:$H${LAST}),0)"
dash["H12"].number_format = "0%"
dash["G13"] = "Payroll rows checked"
dash["H13"] = f'=SUMPRODUCT((Payroll_Input!$I$2:$I${LAST}<>"")*1)'
dash["G14"] = "Payroll rows flagged"
dash["H14"] = f'=SUMPRODUCT((Payroll_Input!$I$2:$I${LAST}<>"OK")*(Payroll_Input!$I$2:$I${LAST}<>""))'
dash["G15"] = "Payroll error rate"
dash["H15"] = "=IFERROR(H14/H13,0)"
dash["H15"].number_format = "0.0%"
for rr in range(12, 16):
    dash.cell(row=rr, column=7).font = BASE
    dash.cell(row=rr, column=8).font = BOLD
    dash.cell(row=rr, column=7).border = BORDER
    dash.cell(row=rr, column=8).border = BORDER
    dash.cell(row=rr, column=8).alignment = Alignment(horizontal="right")

# Charts
bar = BarChart()
bar.type, bar.title = "col", "Hiring Pipeline"
bar.add_data(Reference(dash, min_col=5, min_row=4, max_row=10), titles_from_data=True)
bar.set_categories(Reference(dash, min_col=4, min_row=5, max_row=10))
bar.legend = None
bar.height, bar.width = 7.5, 13
dash.add_chart(bar, "A18")

line = LineChart()
line.title = "Avg Attendance % by Month"
line.add_data(Reference(dash, min_col=8, min_row=4, max_row=7), titles_from_data=True)
line.set_categories(Reference(dash, min_col=7, min_row=5, max_row=7))
line.legend = None
line.y_axis.number_format = "0%"
line.height, line.width = 7.5, 13
dash.add_chart(line, "E18")

# order sheets
order = ["Dashboard", "README", "Employees", "Candidates", "Onboarding", "Attendance", "Payroll_Input", "Settings"]
wb._sheets = [wb[n] for n in order]
wb.save("HR_Dashboard.xlsx")
print("Saved HR_Dashboard.xlsx")
